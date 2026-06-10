"""Convert markdown CV/cover letter to Word (.docx) format."""

from __future__ import annotations

import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

_FONT_BODY = "Calibri"
_FONT_HEADING = "Calibri"
_PT_BODY = 10.5
_PT_H1 = 16
_PT_H2 = 12
_PT_H3 = 11
_PT_CONTACT = 9
_COLOUR_HEADING = "1F2937"  # dark charcoal


def markdown_to_docx(md_content: str, docx_path: str | Path) -> bool:
    """Convert CV/cover letter markdown to a Word document.

    Returns True on success, False on error.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        logger.warning("python-docx not installed — skipping .docx generation")
        return False

    docx_path = Path(docx_path)
    docx_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        doc = Document()

        _set_margins(doc, Cm)
        _configure_styles(doc, Pt, RGBColor)

        lines = md_content.split("\n")
        i = 0
        first_h2_seen = False

        while i < len(lines):
            line = lines[i]
            stripped = line.strip()
            i += 1

            if not stripped or stripped in ("---", "***", "___"):
                continue

            # --- headings ---
            if stripped.startswith("# ") and not stripped.startswith("## "):
                text = stripped[2:].strip()
                p = doc.add_paragraph(style="CV Heading 1")
                _add_rich_runs(p, text, Pt(_PT_H1), RGBColor, bold=True)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue

            if stripped.startswith("## "):
                text = stripped[3:].strip()
                if not first_h2_seen and _is_name_heading(text, md_content):
                    first_h2_seen = True
                    p = doc.add_paragraph(style="CV Heading 1")
                    _add_rich_runs(p, text, Pt(_PT_H1), RGBColor, bold=True)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    first_h2_seen = True
                    p = doc.add_paragraph(style="CV Heading 2")
                    _add_rich_runs(p, text, Pt(_PT_H2), RGBColor, bold=True)
                    _add_bottom_border(p, qn)
                continue

            if stripped.startswith("### "):
                text = stripped[4:].strip()
                p = doc.add_paragraph(style="CV Heading 3")
                _add_rich_runs(p, text, Pt(_PT_H3), RGBColor, bold=True)
                continue

            # --- bullet points ---
            bullet_match = re.match(r"^[-*]\s+(.+)", stripped)
            if bullet_match:
                text = bullet_match.group(1)
                p = doc.add_paragraph(style="CV Bullet")
                _add_bullet_char(p, Pt, RGBColor)
                _add_rich_runs(p, text, Pt(_PT_BODY), RGBColor)
                continue

            # --- contact line (centered bold small text) ---
            if stripped.startswith("**") and (
                "email" in stripped.lower()
                or "phone" in stripped.lower()
                or "@" in stripped
            ):
                p = doc.add_paragraph(style="CV Contact")
                _add_rich_runs(p, stripped, Pt(_PT_CONTACT), RGBColor)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue

            # --- regular paragraph ---
            p = doc.add_paragraph(style="CV Body")
            _add_rich_runs(p, stripped, Pt(_PT_BODY), RGBColor)

        doc.save(str(docx_path))
        return True
    except Exception as exc:
        logger.warning("DOCX generation failed for %s: %s", docx_path, exc)
        return False


# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------

def _set_margins(doc, Cm):
    """Set page margins to a professional CV layout."""
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)


def _configure_styles(doc, Pt, RGBColor):
    """Create custom paragraph styles for CV formatting."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    normal = doc.styles["Normal"]
    normal.font.name = _FONT_BODY
    normal.font.size = Pt(_PT_BODY)
    normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)

    _make = lambda name: doc.styles.add_style(name, 1)  # 1 = WD_STYLE_TYPE.PARAGRAPH

    h1 = _make("CV Heading 1")
    h1.font.name = _FONT_HEADING
    h1.font.size = Pt(_PT_H1)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(_COLOUR_HEADING)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(2)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    h2 = _make("CV Heading 2")
    h2.font.name = _FONT_HEADING
    h2.font.size = Pt(_PT_H2)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(_COLOUR_HEADING)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)

    h3 = _make("CV Heading 3")
    h3.font.name = _FONT_HEADING
    h3.font.size = Pt(_PT_H3)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(_COLOUR_HEADING)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(2)

    body = _make("CV Body")
    body.font.name = _FONT_BODY
    body.font.size = Pt(_PT_BODY)
    body.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    body.paragraph_format.space_before = Pt(1)
    body.paragraph_format.space_after = Pt(3)

    bullet = _make("CV Bullet")
    bullet.font.name = _FONT_BODY
    bullet.font.size = Pt(_PT_BODY)
    bullet.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    bullet.paragraph_format.space_before = Pt(1)
    bullet.paragraph_format.space_after = Pt(2)
    bullet.paragraph_format.left_indent = Pt(18)

    contact = _make("CV Contact")
    contact.font.name = _FONT_BODY
    contact.font.size = Pt(_PT_CONTACT)
    contact.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    contact.paragraph_format.space_before = Pt(0)
    contact.paragraph_format.space_after = Pt(6)
    contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_bottom_border(paragraph, qn):
    """Add a thin bottom border under a section heading."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(
        qn("w:bottom"),
        {
            qn("w:val"): "single",
            qn("w:sz"): "4",
            qn("w:space"): "1",
            qn("w:color"): "9CA3AF",
        },
    )
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_bullet_char(paragraph, Pt, RGBColor):
    """Insert a bullet character as text (avoids font-dependent List Bullet style)."""
    run = paragraph.add_run("\u2022  ")
    run.font.name = _FONT_BODY
    run.font.size = Pt(_PT_BODY)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


# ---------------------------------------------------------------------------
# Rich text: preserves **bold** and *italic* inside paragraphs
# ---------------------------------------------------------------------------

_RICH_RE = re.compile(
    r"(\*\*\*(.+?)\*\*\*"  # ***bold italic***
    r"|\*\*(.+?)\*\*"       # **bold**
    r"|\*(.+?)\*"            # *italic*
    r"|\[([^\]]+)\]\([^)]+\)"  # [link text](url)
    r")"
)


def _add_rich_runs(paragraph, text: str, size, RGBColor, bold: bool = False):
    """Parse markdown inline formatting and add runs with proper styles."""
    pos = 0
    for m in _RICH_RE.finditer(text):
        if m.start() > pos:
            _add_run(paragraph, text[pos : m.start()], size, RGBColor, bold=bold)

        if m.group(2):  # bold+italic
            _add_run(paragraph, m.group(2), size, RGBColor, bold=True, italic=True)
        elif m.group(3):  # bold
            _add_run(paragraph, m.group(3), size, RGBColor, bold=True)
        elif m.group(4):  # italic
            _add_run(paragraph, m.group(4), size, RGBColor, italic=True)
        elif m.group(5):  # link text
            _add_run(paragraph, m.group(5), size, RGBColor, bold=bold)

        pos = m.end()

    if pos < len(text):
        _add_run(paragraph, text[pos:], size, RGBColor, bold=bold)


def _add_run(paragraph, text: str, size, RGBColor, bold=False, italic=False):
    run = paragraph.add_run(text)
    run.font.name = _FONT_BODY
    run.font.size = size
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    if bold:
        run.bold = True
    if italic:
        run.italic = True


# ---------------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------------

def _is_name_heading(text: str, full_content: str) -> bool:
    """Heuristic: the first ## heading in a CV is usually the candidate's name."""
    first_h2 = re.search(r"^## (.+)$", full_content, re.MULTILINE)
    if first_h2:
        clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", first_h2.group(1)).strip()
        return clean == re.sub(r"\*\*([^*]+)\*\*", r"\1", text).strip()
    return False
